#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""在《智能与软件-大模型辅助的程序验证.pptx》研究背景部分插入两页 QCP 示例页。
风格与现有页对齐：标题 32pt 加粗微软雅黑（主题色）、蓝色小节条（003CB4）、
正文微软雅黑、强调色 C00000。不修改任何已有页面。"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = '智能与软件-大模型辅助的程序验证.pptx'

prs = Presentation(SRC)
slides = list(prs.slides)
s_content = slides[8]           # 研究一 内容页，提供标题与小节条样式
layout = slides[4].slide_layout

DARK = RGBColor(0x44, 0x54, 0x6A)     # 主题正文色 dk1
BLUE = RGBColor(0x00, 0x3C, 0xB4)     # 模板小节条蓝
RED = RGBColor(0xC0, 0x00, 0x00)      # 模板强调红
CODEBG = RGBColor(0xF5, 0xF5, 0xF5)
BORDER = RGBColor(0xB8, 0xC4, 0xD8)

# 样式原型
proto_title = proto_label = None
for sh in s_content.shapes:
    if sh.name == '文本框 1':
        proto_title = sh
    if sh.name == '矩形 8':
        proto_label = sh
assert proto_title is not None and proto_label is not None


def clone_shape(slide, proto):
    el = copy.deepcopy(proto._element)
    slide.shapes._spTree.append(el)
    from pptx.shapes.autoshape import Shape
    return Shape(el, slide.shapes)


def set_shape_text(shape, text):
    """替换克隆形状的文字，保留首个 run 的样式。"""
    tf = shape.text_frame
    paras = tf.paragraphs
    # 删除多余段落
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    runs = p0.runs
    for r in list(runs[1:]):
        r._r.getparent().remove(r._r)
    runs[0].text = text


def make_run(p, text, size, bold=False, color=DARK, mono=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea'):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set('typeface', 'Consolas' if mono else '微软雅黑')
    return r


def add_text(slide, x, y, w, h, items, anchor_top=True):
    """items: [(segments, size, level, bullet, space_after)]，segments 为 [(text,{bold,color,mono})]"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for segments, size, level, bullet, spa in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(spa)
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '●'}))
        for text, opt in segments:
            make_run(p, text, size, bold=opt.get('bold', False),
                     color=opt.get('color', DARK), mono=opt.get('mono', False))
    return tb


def add_label(slide, x, y, text, w=Inches(3.4)):
    lab = clone_shape(slide, proto_label)
    lab.left, lab.top, lab.width = x, y, w
    lab.height = proto_label.height
    set_shape_text(lab, text)
    return lab


def add_code(slide, x, y, w, h, lines, size=12):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = BORDER; box.line.width = Pt(0.75)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12); tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.08); tf.margin_bottom = Inches(0.06)
    first = True
    for ln, color in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(0)
        make_run(p, ln if ln else ' ', size, color=color, mono=True)
    return box


def add_arrow(slide, x, y):
    ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, Inches(0.4), Inches(0.3))
    ar.fill.solid(); ar.fill.fore_color.rgb = BLUE; ar.line.fill.background()


M = Inches(0.45)
CODE_C = DARK
ANN_C = BLUE

# ============================================================
# 新页 A：例子页（程序 + 非形式需求，完整规约，取自 OUTPUT/xizi）
# ============================================================
sA = prs.slides.add_slide(layout)
tA = clone_shape(sA, proto_title)
set_shape_text(tA, '研究背景：一个例子——单链表求长程序的验证')

add_label(sA, M, Inches(1.2), '非形式需求', w=Inches(2.2))
add_text(sA, M, Inches(1.75), Inches(3.55), Inches(2.7), [
    ([('输入：单链表头指针 linklist', {})], 13, 0, True, 8),
    ([('输出：链表的结点个数', {})], 13, 0, True, 8),
    ([('约束：不修改链表（含各结点 payload）', {})], 13, 0, True, 8),
    ([('空表返回 0', {})], 13, 0, True, 8),
])
add_label(sA, M, Inches(4.35), '案例来源', w=Inches(2.2))
add_text(sA, M, Inches(4.9), Inches(3.55), Inches(1.9), [
    ([('西子实时操作系统单链表模块', {})], 13, 0, True, 6),
    ([('已完成端到端验证的真实案例', {})], 13, 0, True, 6),
    ([('标注为归档文件中的完整规约', {'bold': True, 'color': RED})], 13, 0, True, 6),
])

code = [
    ('/*@ Extern Coq (xizi_sll_head_payload : ...) */', ANN_C),
    ('unsigned int SingleLinkListGetLen(const SysSingleLinklistType *linklist)', CODE_C),
    ('/*@ With {A} (storeA : Z -> A -> Assertion) (l : list (sll_payload_node A))', ANN_C),
    ('    Require  xizi_sll_head_payload(storeA, linklist, l)', ANN_C),
    ('    Ensure   __return == Zlength(l) &&', ANN_C),
    ('             xizi_sll_head_payload(storeA, linklist, l) */', ANN_C),
    ('{', CODE_C),
    ('    /*@ Assert', ANN_C),
    ('          linklist == linklist@pre &&', ANN_C),
    ('          xizi_sll_head(linklist, xizi_sll_payload_addresses(l)) *', ANN_C),
    ('          xizi_sll_remove_payloads(storeA, l) */', ANN_C),
    ('    unsigned int length = 0;', CODE_C),
    ('    const SysSingleLinklistType *tmp_list = linklist->node_next;', CODE_C),
    ('    /*@ Inv exists first l1 l2,', ANN_C),
    ('          xizi_sll_payload_addresses(l) == app(l1, l2) &&', ANN_C),
    ('          length == Zlength(l1) &&', ANN_C),
    ('          linklist == linklist@pre && linklist != 0 &&', ANN_C),
    ('          linklist->node_next == first &&', ANN_C),
    ('          xizi_sllseg(first, tmp_list, l1) *', ANN_C),
    ('          xizi_sll(tmp_list, l2) *', ANN_C),
    ('          xizi_sll_remove_payloads(storeA, l) */', ANN_C),
    ('    while (tmp_list != (void *)0) {', CODE_C),
    ('        tmp_list = tmp_list->node_next;', CODE_C),
    ('        length++;', CODE_C),
    ('    }', CODE_C),
    ('    return length;', CODE_C),
    ('}', CODE_C),
]
add_label(sA, Inches(4.3), Inches(1.2), '带 QCP 标注的程序（含循环）', w=Inches(4.6))
add_code(sA, Inches(4.3), Inches(1.75), Inches(8.55), Inches(4.75), code, size=10.5)

add_text(sA, Inches(4.3), Inches(6.6), Inches(8.55), Inches(0.8), [
    ([('Require / Ensure', {'bold': True, 'color': RED}), ('：前置/后置条件（功能规约）；', {}),
      ('Assert / Inv', {'bold': True, 'color': RED}), ('：中间辅助断言（循环不变式）；', {}),
      ('sll 系列谓词', {'bold': True, 'color': RED}), ('：内存中链表与数学列表的对应。', {})], 12, 0, False, 4),
])

# ============================================================
# 新页 B：验证的三个组成部分（三行式布局，与 CAV-Bench 对应）
# ============================================================
sB = prs.slides.add_slide(layout)
tB = clone_shape(sB, proto_title)
set_shape_text(tB, '研究背景：程序验证的三个组成部分')

# 顶部流程
flow = ['带标注的 C 程序', 'QCP 符号执行生成验证条件', 'Rocq 证明', '验证通过']
fx, fw, fgap = M, Inches(2.75), Inches(0.42)
for i, t in enumerate(flow):
    box = sB.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx, Inches(1.12), fw, Inches(0.62))
    box.adjustments[0] = 0.14
    box.fill.solid(); box.fill.fore_color.rgb = BLUE
    box.line.fill.background()
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    make_run(p, t, 13, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    if i < 3:
        add_arrow(sB, fx + fw + Inches(0.01), Inches(1.28))
    fx += fw + fgap

# 三行：每行左侧蓝色标签块 + 右侧内容
rows = [
    ('① 功能规约', '阶段 I：规约生成', [
        ([('Require  xizi_sll_head_payload(storeA, linklist, l)', {'mono': True})], 11.5, 0, False, 3),
        ([('Ensure   __return == Zlength(l) && xizi_sll_head_payload(...)', {'mono': True})], 11.5, 0, False, 6),
        ([('含义：输入链表内容恰为数学列表 l；返回 l 的长度；链表不被修改。', {})], 12, 0, False, 2),
    ]),
    ('② 中间辅助断言', '阶段 II：辅助标注生成', [
        ([('循环不变式 Inv：已扫描段 l1 已计数 length == Zlength(l1)，剩余段 l2 仍为链表，地址序列满足 app(l1, l2) 不变', {})], 12, 0, False, 5),
        ([('循环前 Assert：将头结点拆解为“地址序列 + 载荷分离”的形式，为符号执行展开内存布局', {})], 12, 0, False, 5),
        ([('循环与指针程序验证的关键，也是学术难点。', {'color': RED})], 12, 0, False, 2),
    ]),
    ('③ 验证条件生成与证明', '阶段 III：证明生成', [
        ([('QCP 符号执行沿程序路径自动生成验证条件：内存安全 safety VC + 功能正确性 return VC', {})], 12, 0, False, 5),
        ([('在 Rocq 中完成证明，coqc 检查通过即验证完成。', {})], 12, 0, False, 2),
    ]),
]
ry = Inches(2.0)
for head, sub, body in rows:
    lab = sB.shapes.add_shape(MSO_SHAPE.RECTANGLE, M, ry, Inches(2.9), Inches(1.28))
    lab.fill.solid(); lab.fill.fore_color.rgb = BLUE; lab.line.fill.background()
    tf = lab.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    make_run(p, head, 14, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    make_run(p2, sub, 12, bold=False, color=RGBColor(0xFF, 0xFF, 0xFF))
    add_text(sB, Inches(3.65), ry + Inches(0.06), Inches(9.2), Inches(1.3), body)
    ry += Inches(1.52)

add_text(sB, M, Inches(6.6), Inches(12.4), Inches(0.7), [
    ([('CAV-Bench 三种任务设置即按涵盖阶段划分：', {'bold': True}),
      ('I+II+III（端到端）、II+III、III（仅证明）', {'bold': True, 'color': RED})], 14, 0, False, 2),
])

# ============================================================
# 两页移动到第 5 页之后（索引 5、6）
# ============================================================
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
newA, newB = ids[-2], ids[-1]
sldIdLst.remove(newA); sldIdLst.remove(newB)
sldIdLst.insert(5, newA)
sldIdLst.insert(6, newB)

prs.save(SRC)
print('done, total slides =', len(prs.slides._sldIdLst))
