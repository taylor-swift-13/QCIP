# -*- coding: utf-8 -*-
"""基于季度汇报模板生成独立的三页演示文稿。

原始季度汇报只作为主题与版式来源，不会被修改。
"""

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "2026年第二季度工作总结汇报.pptx"
OUTPUT = HERE / "自动机生成与规约验证测试_三页汇报.pptx"
BLUE = RGBColor(0x00, 0x3C, 0xB4)
BLUE_2 = RGBColor(0x3C, 0x56, 0xE5)
TEAL = RGBColor(0x07, 0xAC, 0xBD)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
ORANGE = RGBColor(0xE7, 0x8A, 0x20)
RED = RGBColor(0xC0, 0x00, 0x00)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x65, 0x73)
LIGHT_BLUE = RGBColor(0xEC, 0xF3, 0xFF)
LIGHT_TEAL = RGBColor(0xE8, 0xF8, 0xFA)
LIGHT_ORANGE = RGBColor(0xFF, 0xF4, 0xE5)
LIGHT_GRAY = RGBColor(0xF3, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"
MONO = "Consolas"


def remove_all_slides(prs):
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        rel_id = slide_id.rId
        slide_ids.remove(slide_id)
        prs.part.drop_rel(rel_id)


def clear_slide(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def add_cloned_shape(slide, element):
    slide.shapes._spTree.append(deepcopy(element))


def title_slide(prs, layout, title_template, number_template, title):
    slide = prs.slides.add_slide(layout)
    clear_slide(slide)
    add_cloned_shape(slide, title_template)
    add_cloned_shape(slide, number_template)
    for shape in slide.shapes:
        if shape.name == "标题 1":
            shape.text = title
            for p in shape.text_frame.paragraphs:
                p.font.name = FONT
    return slide


def set_shape_style(shape, fill, line=None, radius=True):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    shape.line.width = Pt(1.15)


def add_box(
    slide,
    x,
    y,
    w,
    h,
    text="",
    *,
    fill=WHITE,
    line=BLUE_2,
    font_size=15,
    color=INK,
    bold=False,
    align=PP_ALIGN.CENTER,
    radius=True,
    margin=0.08,
    font=FONT,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_style(shape, fill, line)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin / 1.3)
    tf.margin_bottom = Inches(margin / 1.3)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    paragraphs,
    *,
    font_size=15,
    color=INK,
    bullet=False,
    line_spacing=1.05,
    space_after=6,
    font=FONT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    for i, item in enumerate(paragraphs):
        if isinstance(item, tuple):
            text, item_color, item_bold = item
        else:
            text, item_color, item_bold = item, color, False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.name = font
        p.font.size = Pt(font_size)
        p.font.color.rgb = item_color
        p.font.bold = item_bold
        p.level = 0
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        if bullet:
            p.text = "• " + p.text
    return box


def add_banner(slide, index, text):
    add_box(
        slide,
        0.55,
        1.00,
        12.20,
        0.48,
        f"{index}/3   {text}",
        fill=LIGHT_BLUE,
        line=BLUE,
        font_size=19,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.LEFT,
        margin=0.13,
    )


def add_chevron(slide, x, y, w=0.25, h=0.32, color=TEAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    return shape


def add_card(
    slide,
    x,
    y,
    w,
    h,
    title,
    lines,
    *,
    accent=BLUE,
    fill=WHITE,
    body_font_size=13,
):
    card = add_box(slide, x, y, w, h, fill=fill, line=accent, radius=True)
    add_box(
        slide,
        x + 0.08,
        y + 0.08,
        w - 0.16,
        0.43,
        title,
        fill=accent,
        line=accent,
        font_size=16,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.LEFT,
        margin=0.12,
    )
    add_text(
        slide,
        x + 0.18,
        y + 0.65,
        w - 0.36,
        h - 0.78,
        lines,
        font_size=body_font_size,
        color=INK,
        bullet=True,
        line_spacing=0.95,
        space_after=3,
    )
    return card


def add_stat(slide, x, y, w, number, label, color=BLUE):
    add_box(
        slide,
        x,
        y,
        w,
        0.67,
        fill=WHITE,
        line=color,
        radius=True,
    )
    add_text(
        slide,
        x + 0.08,
        y + 0.09,
        0.88,
        0.42,
        [(number, color, True)],
        font_size=22,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        x + 0.95,
        y + 0.10,
        w - 1.03,
        0.40,
        [label],
        font_size=12.5,
        color=MUTED,
        valign=MSO_ANCHOR.MIDDLE,
    )


def slide_automation(prs, layout, title_template, number_template):
    slide = title_slide(
        prs, layout, title_template, number_template, "研究进展3：知识模型提取"
    )
    add_banner(slide, 1, "源码到自动机：批量生成可审查的源码级控制流图")

    steps = ["C 源码", "文本级解析", "控制结构建模", "Graphviz 渲染", "PNG + 日志"]
    x = 0.55
    for i, step in enumerate(steps):
        add_box(
            slide,
            x,
            1.74,
            1.95,
            0.63,
            step,
            fill=LIGHT_BLUE if i < 4 else LIGHT_TEAL,
            line=BLUE if i < 4 else TEAL,
            font_size=13,
            color=BLUE if i < 4 else TEAL,
            bold=True,
        )
        if i < len(steps) - 1:
            add_chevron(slide, x + 2.05, 1.90, 0.23, 0.28)
        x += 2.48

    add_card(
        slide,
        0.55,
        2.68,
        5.85,
        1.45,
        "不依赖编译环境",
        ["无需头文件、clang/LLVM", "不完整或暂不可编译的 C 文件也可尽量生成"],
        accent=BLUE,
        fill=LIGHT_BLUE,
    )
    add_card(
        slide,
        6.88,
        2.68,
        5.85,
        1.45,
        "覆盖主要控制结构",
        ["if / else、各类循环、switch / case", "识别 break、continue、return 与 goto 出口"],
        accent=TEAL,
        fill=LIGHT_TEAL,
    )
    add_card(
        slide,
        0.55,
        4.36,
        5.85,
        1.45,
        "可追踪、可审查",
        ["日志记录函数、节点数、边数与节点源码", "终止语句后的不可达代码仍显示，不静默丢语句"],
        accent=GREEN,
        fill=RGBColor(0xEC, 0xF8, 0xF0),
    )
    add_card(
        slide,
        6.88,
        4.36,
        5.85,
        1.45,
        "批处理与离线交付",
        ["Docker 封装，支持内网无网环境", "当前样例集 16/16 成功生成；结果可直接进入知识库"],
        accent=ORANGE,
        fill=LIGHT_ORANGE,
    )

    add_box(
        slide,
        0.55,
        6.09,
        12.18,
        0.55,
        "定位：输出是面向批量知识抽取的源码级近似自动机，强调结构完整、可追踪和可离线运行。",
        fill=LIGHT_GRAY,
        line=RGBColor(0xD0, 0xD6, 0xDE),
        font_size=14,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    return slide


def slide_spec(prs, layout, title_template, number_template):
    slide = title_slide(
        prs, layout, title_template, number_template, "研究进展3：知识模型提取"
    )
    add_banner(slide, 2, "规约与循环不变式：先定义性质，再填写 annotation")

    steps = [
        ("读 C 程序\n明确目标", BLUE),
        ("定义数学\nspec", BLUE_2),
        ("填写 Require /\nEnsure / Inv", TEAL),
        ("QCP 交互\n检查与修正", GREEN),
        ("annotation-\nchecking", ORANGE),
    ]
    x = 0.55
    for i, (text, color) in enumerate(steps):
        add_box(
            slide,
            x,
            1.72,
            2.05,
            0.66,
            text,
            fill=WHITE,
            line=color,
            font_size=14,
            color=color,
            bold=True,
        )
        if i < 4:
            add_chevron(slide, x + 2.16, 1.90, 0.25, 0.28, color=TEAL)
        x += 2.55

    add_card(
        slide,
        0.55,
        2.67,
        3.85,
        2.70,
        "① 先写数学规格（spec-first）",
        [
            "在 case_lib 中定义函数真正要满足的数学性质",
            "描述输入与输出的关系，不在 Rocq 中照抄 C 算法或循环状态机",
        ],
        accent=BLUE,
        fill=LIGHT_BLUE,
    )
    add_box(
        slide,
        0.84,
        4.62,
        3.27,
        0.46,
        "Post(input, output, state′)",
        fill=WHITE,
        line=BLUE,
        font_size=14,
        color=BLUE,
        bold=True,
        font=MONO,
    )

    add_card(
        slide,
        4.74,
        2.67,
        3.85,
        2.70,
        "② 再写 C annotation",
        [
            "Require：变量范围与必要的内存资源",
            "Ensure：调用数学 spec 表达最终结果",
            "Assert / Inv：连接函数调用、分支与循环中间状态",
        ],
        accent=TEAL,
        fill=LIGHT_TEAL,
    )
    add_box(
        slide,
        5.03,
        4.62,
        3.27,
        0.46,
        "Require  →  Inv  →  Ensure",
        fill=WHITE,
        line=TEAL,
        font_size=12.5,
        color=TEAL,
        bold=True,
        font=MONO,
    )

    add_card(
        slide,
        8.93,
        2.67,
        3.85,
        2.70,
        "③ 不变式回答三个问题",
        [
            "初始化：进入循环前是否成立？",
            "保持：执行一次循环体后是否仍成立？",
            "退出：不变式与退出条件能否推出 Ensure？",
        ],
        accent=ORANGE,
        fill=LIGHT_ORANGE,
    )
    add_box(
        slide,
        0.55,
        5.72,
        12.23,
        0.86,
        "QCP 检查跑到文件尾  →  annotation-checking 检查 spec、函数规约和不变式  →  通过后才运行正式符号执行并生成 VC。\n任一步失败，都回到前面修正规约或 annotation，不把问题带入证明阶段。",
        fill=WHITE,
        line=BLUE,
        font_size=14,
        color=INK,
        bold=False,
        align=PP_ALIGN.CENTER,
    )
    return slide


def add_vertical_flow(slide, x, y, w, labels, accent, pale):
    box_h = 0.56
    gap = 0.20
    for i, label in enumerate(labels):
        add_box(
            slide,
            x,
            y + i * (box_h + gap),
            w,
            box_h,
            label,
            fill=pale if i < len(labels) - 1 else WHITE,
            line=accent,
            font_size=13.5,
            color=accent if i < len(labels) - 1 else INK,
            bold=True,
        )
        if i < len(labels) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.DOWN_ARROW,
                Inches(x + w / 2 - 0.10),
                Inches(y + box_h + i * (box_h + gap) + 0.02),
                Inches(0.20),
                Inches(0.16),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = accent
            arrow.line.color.rgb = accent


def slide_validation(prs, layout, title_template, number_template):
    slide = title_slide(
        prs, layout, title_template, number_template, "研究进展3：知识模型提取"
    )
    add_banner(slide, 3, "验证与测试：一套规约的两级可信闭环")

    add_box(
        slide,
        0.55,
        1.70,
        5.90,
        3.85,
        fill=LIGHT_BLUE,
        line=BLUE,
    )
    add_box(
        slide,
        0.77,
        1.86,
        5.46,
        0.47,
        "路线 A｜QCP + Rocq：全称形式化验证（16 个程序）",
        fill=BLUE,
        line=BLUE,
        font_size=16,
        color=WHITE,
        bold=True,
    )
    add_vertical_flow(
        slide,
        1.27,
        2.56,
        4.46,
        [
            "带 Require / Ensure / Inv 的 C 程序",
            "符号执行自动生成验证条件（VC）",
            "Rocq 完成证明并由内核检查",
            "结论：对所有满足前置条件的输入成立",
        ],
        BLUE,
        WHITE,
    )
    add_box(
        slide,
        6.87,
        1.70,
        5.90,
        3.85,
        fill=LIGHT_TEAL,
        line=TEAL,
    )
    add_box(
        slide,
        7.09,
        1.86,
        5.46,
        0.47,
        "路线 B｜FloatTest：位级差分规约测试（79 个程序）",
        fill=TEAL,
        line=TEAL,
        font_size=16,
        color=WHITE,
        bold=True,
    )
    add_vertical_flow(
        slide,
        7.59,
        2.56,
        4.46,
        [
            "同一输入分别执行 C 实现与 Rocq 可执行规约",
            "按 IEEE 754 bit pattern 逐位比较输出",
            "每个测试向量生成一条可编译 Rocq 定理",
            "结论：在采样输入上逐比特一致（约 7.87 万条定理）",
        ],
        TEAL,
        WHITE,
    )
    add_box(
        slide,
        0.55,
        5.83,
        12.22,
        0.83,
        fill=WHITE,
        line=ORANGE,
    )
    add_box(
        slide,
        0.76,
        6.01,
        1.72,
        0.46,
        "不是二选一",
        fill=ORANGE,
        line=ORANGE,
        font_size=15,
        color=WHITE,
        bold=True,
    )
    add_text(
        slide,
        2.72,
        5.96,
        9.72,
        0.58,
        [
            "证明可达的程序做全称证明；浮点／libm 密集、暂难经济证明的程序先做位级定理化测试；任何失败都反馈到规约与不变式提取环节。"
        ],
        font_size=14,
        color=INK,
        valign=MSO_ANCHOR.MIDDLE,
    )
    return slide


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    prs = Presentation(SOURCE)
    reference = prs.slides[9]
    layout = reference.slide_layout
    title_template = deepcopy(
        next(shape._element for shape in reference.shapes if shape.name == "标题 1")
    )
    number_template = deepcopy(
        next(
            shape._element
            for shape in reference.shapes
            if shape.name == "灯片编号占位符 1"
        )
    )

    remove_all_slides(prs)
    slide_automation(prs, layout, title_template, number_template)
    slide_spec(prs, layout, title_template, number_template)
    slide_validation(prs, layout, title_template, number_template)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
